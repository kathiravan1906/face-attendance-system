import cv2
import os
from deepface import DeepFace
from datetime import datetime
from openpyxl import Workbook, load_workbook

# Image folder
path = 'images'
known_faces = os.listdir(path)

# Excel file setup
excel_file = "attendance.xlsx"

if not os.path.exists(excel_file):
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Date", "Time"])
    wb.save(excel_file)

# Function to mark attendance
def markAttendance(name):

    wb = load_workbook(excel_file)
    ws = wb.active

    now = datetime.now()
    date = now.strftime("%Y-%m-%d")
    time = now.strftime("%H:%M:%S")

    already_present = False

    # Check existing attendance
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == name and row[1] == date:
            already_present = True
            break

    if not already_present:
        ws.append([name, date, time])
        wb.save(excel_file)
        return "PRESENT"

    else:
        return "ALREADY DETECTED"

# Start webcam
cap = cv2.VideoCapture(0)

while True:
    ret, frame = cap.read()

    recognized = False

    for img_name in known_faces:

        img_path = os.path.join(path, img_name)

        try:
            result = DeepFace.verify(
                frame,
                img_path,
                enforce_detection=False
            )

            if result['verified']:

                name = os.path.splitext(img_name)[0]

                status = markAttendance(name)

                now = datetime.now()
                current_time = now.strftime("%H:%M:%S")

                # Green color for present
                if status == "PRESENT":

                    cv2.putText(frame,
                                f"{name.upper()} - PRESENT",
                                (50,50),
                                cv2.FONT_HERSHEY_SIMPLEX,
                                1,
                                (0,255,0),
                                2)

                # Yellow for already detected
                else:

                    cv2.putText(frame,
                                f"{name.upper()} - ALREADY DETECTED",
                                (50,50),
                                cv2.FONT_HERSHEY_SIMPLEX,
                                1,
                                (0,255,255),
                                2)

                # Show time
                cv2.putText(frame,
                            f"Time: {current_time}",
                            (50,100),
                            cv2.FONT_HERSHEY_SIMPLEX,
                            1,
                            (255,255,255),
                            2)

                recognized = True
                break

        except:
            pass

    # Unknown person
    if not recognized:

        cv2.putText(frame,
                    "UNKNOWN PERSON",
                    (50,50),
                    cv2.FONT_HERSHEY_SIMPLEX,
                    1,
                    (0,0,255),
                    2)

    cv2.imshow("Face Attendance System", frame)

    # Press q to exit
    if cv2.waitKey(1) & 0xFF == ord('q'):
        break

cap.release()
cv2.destroyAllWindows()
